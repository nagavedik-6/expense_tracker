import csv
import io
from datetime import timedelta
from decimal import Decimal

from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.db.models import Sum
from django.utils import timezone

from expenses.models import Expense
from income.models import Income

# PDF generation
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

# Excel generation
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def _get_report_data(user, period, start_date=None, end_date=None):
    """Get expenses and income for a period."""
    today = timezone.now().date()

    if period == 'daily':
        start_date = start_date or today
        end_date = end_date or today
    elif period == 'weekly':
        start_date = start_date or (today - timedelta(days=today.weekday()))
        end_date = end_date or (start_date + timedelta(days=6))
    elif period == 'monthly':
        start_date = start_date or today.replace(day=1)
        end_date = end_date or today
    elif period == 'yearly':
        start_date = start_date or today.replace(month=1, day=1)
        end_date = end_date or today

    expenses = Expense.objects.filter(
        user=user, date__gte=start_date, date__lte=end_date
    ).select_related('category').order_by('-date')

    incomes = Income.objects.filter(
        user=user, date__gte=start_date, date__lte=end_date
    ).order_by('-date')

    total_expenses = expenses.aggregate(total=Sum('amount'))['total'] or Decimal('0')
    total_income = incomes.aggregate(total=Sum('amount'))['total'] or Decimal('0')

    # Category breakdown
    category_data = expenses.values('category__name').annotate(
        total=Sum('amount')).order_by('-total')

    return {
        'expenses': expenses,
        'incomes': incomes,
        'total_expenses': total_expenses,
        'total_income': total_income,
        'net_savings': total_income - total_expenses,
        'category_data': category_data,
        'start_date': start_date,
        'end_date': end_date,
        'period': period,
    }


@login_required
def report_dashboard(request):
    """Report dashboard with period selection."""
    period = request.GET.get('period', 'monthly')
    start_date = request.GET.get('start_date')
    end_date = request.GET.get('end_date')

    if start_date:
        from datetime import date as dt_date
        start_date = dt_date.fromisoformat(start_date)
    if end_date:
        from datetime import date as dt_date
        end_date = dt_date.fromisoformat(end_date)

    data = _get_report_data(request.user, period, start_date, end_date)
    currency_symbol = request.user.profile.currency_symbol if hasattr(request.user, 'profile') else '₹'
    data['currency_symbol'] = currency_symbol

    return render(request, 'reports/report_dashboard.html', data)


@login_required
def export_csv(request):
    """Export expenses as CSV."""
    period = request.GET.get('period', 'monthly')
    data = _get_report_data(request.user, period)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="expense_report_{period}.csv"'

    writer = csv.writer(response)
    writer.writerow(['Date', 'Title', 'Category', 'Amount', 'Payment Method', 'Notes'])
    for exp in data['expenses']:
        writer.writerow([
            exp.date, exp.title,
            exp.category.name if exp.category else 'N/A',
            exp.amount, exp.get_payment_method_display(),
            exp.notes or ''
        ])

    writer.writerow([])
    writer.writerow(['', '', '', f'Total: {data["total_expenses"]}', '', ''])

    return response


@login_required
def export_excel(request):
    """Export expenses as Excel."""
    period = request.GET.get('period', 'monthly')
    data = _get_report_data(request.user, period)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Expense Report ({period})"

    # Header style
    header_fill = PatternFill(start_color='1a73e8', end_color='1a73e8', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=12)

    headers = ['Date', 'Title', 'Category', 'Amount', 'Payment Method', 'Notes']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    for row, exp in enumerate(data['expenses'], 2):
        ws.cell(row=row, column=1, value=str(exp.date))
        ws.cell(row=row, column=2, value=exp.title)
        ws.cell(row=row, column=3, value=exp.category.name if exp.category else 'N/A')
        ws.cell(row=row, column=4, value=float(exp.amount))
        ws.cell(row=row, column=5, value=exp.get_payment_method_display())
        ws.cell(row=row, column=6, value=exp.notes or '')

    # Summary row
    summary_row = len(data['expenses']) + 3
    ws.cell(row=summary_row, column=3, value='Total Expenses:').font = Font(bold=True)
    ws.cell(row=summary_row, column=4, value=float(data['total_expenses'])).font = Font(bold=True)
    ws.cell(row=summary_row + 1, column=3, value='Total Income:').font = Font(bold=True)
    ws.cell(row=summary_row + 1, column=4, value=float(data['total_income'])).font = Font(bold=True)
    ws.cell(row=summary_row + 2, column=3, value='Net Savings:').font = Font(bold=True)
    ws.cell(row=summary_row + 2, column=4, value=float(data['net_savings'])).font = Font(bold=True)

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col) + 2
        ws.column_dimensions[col[0].column_letter].width = max_len

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="expense_report_{period}.xlsx"'
    wb.save(response)
    return response


@login_required
def export_pdf(request):
    """Export expenses as PDF."""
    period = request.GET.get('period', 'monthly')
    data = _get_report_data(request.user, period)
    currency_symbol = request.user.profile.currency_symbol if hasattr(request.user, 'profile') else '₹'

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="expense_report_{period}.pdf"'

    doc = SimpleDocTemplate(response, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    # Title
    title = Paragraph(f"Expense Report — {period.title()}", styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 0.3 * inch))

    subtitle = Paragraph(
        f"Period: {data['start_date']} to {data['end_date']}", styles['Normal']
    )
    elements.append(subtitle)
    elements.append(Spacer(1, 0.3 * inch))

    # Summary
    summary_data = [
        ['Total Income', f'{currency_symbol}{data["total_income"]}'],
        ['Total Expenses', f'{currency_symbol}{data["total_expenses"]}'],
        ['Net Savings', f'{currency_symbol}{data["net_savings"]}'],
    ]
    summary_table = Table(summary_data, colWidths=[3 * inch, 2 * inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8f9fa')),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#333333')),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.4 * inch))

    # Expense table
    if data['expenses']:
        elements.append(Paragraph("Expense Details", styles['Heading2']))
        elements.append(Spacer(1, 0.2 * inch))

        table_data = [['Date', 'Title', 'Category', 'Amount', 'Payment']]
        for exp in data['expenses']:
            table_data.append([
                str(exp.date),
                exp.title[:30],
                (exp.category.name if exp.category else 'N/A')[:20],
                f'{currency_symbol}{exp.amount}',
                exp.get_payment_method_display()
            ])

        exp_table = Table(table_data, colWidths=[1.1*inch, 1.8*inch, 1.3*inch, 1*inch, 1.2*inch])
        exp_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a73e8')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f0f4ff')]),
        ]))
        elements.append(exp_table)

    doc.build(elements)
    return response
